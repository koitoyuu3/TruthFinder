#!/usr/bin/env python3
"""Fix content issues in 5 variable-description paragraphs: remove variables not in their formula."""

import copy
from lxml import etree
from docx import Document

W = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
M = 'http://schemas.openxmlformats.org/officeDocument/2006/math'


def qw(tag):
    return f'{{{W}}}{tag}'


def find_wt(element):
    """Find the first w:t element in a w:ins or w:r."""
    for wt in element.iter(qw('t')):
        return wt
    return None


def show(para, label):
    children = list(para._element)
    print(f'\n{label}:')
    for ci, child in enumerate(children):
        tag = child.tag.split('}')[1] if '}' in child.tag else child.tag
        ns = 'w' if W in child.tag else ('m' if M in child.tag else '?')
        if ns == 'w' and tag == 'ins':
            texts = [t.text or '' for t in child.iter(f'{{{W}}}t')]
            print(f'  [{ci}] w:ins: "{" ".join(texts).strip()[:80]}"')
        elif ns == 'm' and tag == 'oMath':
            texts = [t.text or '' for t in child.iter(f'{{{M}}}t')]
            print(f'  [{ci}] m:oMath: "{" ".join(texts).strip()[:80]}"')
        elif ns == 'w' and tag == 'pPr':
            print(f'  [{ci}] w:pPr')


def main():
    src = '/docs/【发明】【北京邮电大学】专利书-LASO.docx'
    doc = Document(src)

    # ── P78 (式4): Remove σ_o full definition (from 式3), keep reference ──
    print('=== Fixing P78 (式4) ===')
    p78 = doc.paragraphs[78]
    children = list(p78._element)
    # [2] m:oMath σ_o  ← keep
    # [3] w:ins "为鲁棒尺度估计，定义为" ← change text
    # [4] m:oMath σ_o = δ_o / Φ^{-1}... ← remove
    # [5] w:ins "；" ← remove
    # [6] m:oMath Φ^{-1}(0.75) ← remove
    # [7] w:ins "≈ 0.67449..." ← remove

    wt3 = find_wt(children[3])
    wt3.text = '为式（3）所定义的鲁棒尺度；'
    wt3.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')

    for idx in [7, 6, 5, 4]:  # remove in reverse order
        p78._element.remove(children[idx])

    show(p78, 'P78 after fix')

    # ── P81 (式5): Remove σ_o / ε_tol convergence note ──
    print('\n=== Fixing P81 (式5) ===')
    p81 = doc.paragraphs[81]
    children = list(p81._element)
    # [15] w:ins "个保留源的数值报告。该过程重复迭代...；当" ← truncate
    # [16] m:oMath σ_o ← remove
    # [17] w:ins "低于容差阈值" ← remove
    # [18] m:oMath ε_tol ← remove
    # [19] w:ins "时，直接返回..." ← remove

    wt15 = find_wt(children[15])
    wt15.text = '个保留源的数值报告。'

    for idx in [19, 18, 17, 16]:
        p81._element.remove(children[idx])

    show(p81, 'P81 after fix')

    # ── P87 (式7): Remove sim_o^{i,j} definition ──
    print('\n=== Fixing P87 (式7) ===')
    p87 = doc.paragraphs[87]
    children = list(p87._element)
    # [15] w:ins "为 L2 范数...两条文本声明之间的语义相似度定义为其语义向量的余弦相似度" ← truncate
    # [16] m:oMath sim_o^{i,j} = ... ← remove
    # [17] w:ins "。" ← remove

    wt15 = find_wt(children[15])
    wt15.text = '为 L2 范数，用于将嵌入向量归一化为单位长度。'

    for idx in [17, 16]:
        p87._element.remove(children[idx])

    show(p87, 'P87 after fix')

    # ── P89 (式8): Remove L_o^i, L_o^j, π_num ──
    print('\n=== Fixing P89 (式8) ===')
    p89 = doc.paragraphs[89]
    children = list(p89._element)
    # [17] w:ins "；" ← change to "。"
    # [18]-[27] ← remove all

    wt17 = find_wt(children[17])
    wt17.text = '。'

    for idx in range(27, 17, -1):  # 27 down to 18
        p89._element.remove(children[idx])

    show(p89, 'P89 after fix')

    # ── P100 (式10): Remove is_valid ──
    print('\n=== Fixing P100 (式10) ===')
    p100 = doc.paragraphs[100]
    children = list(p100._element)
    # [9] w:ins "为源可靠性更新电路...；∧ 表示逻辑合取，即电路输出有效信号" ← modify
    # [10] m:oMath is_valid = 1 ← remove
    # [11] w:ins "当且仅当三个子模块的约束全部成立。" ← remove

    wt9 = find_wt(children[9])
    wt9.text = '为源可靠性更新电路，基于通道级一致性得分实现 Beta 后验参数更新；∧ 表示逻辑合取，即整体电路约束成立当且仅当三个子模块的约束全部成立。'
    wt9.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')

    for idx in [11, 10]:
        p100._element.remove(children[idx])

    show(p100, 'P100 after fix')

    # Save
    doc.save(src)
    print(f'\nSaved to: {src}')


if __name__ == '__main__':
    main()
