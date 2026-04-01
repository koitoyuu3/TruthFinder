from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


LINK_RE = re.compile(r"\[([^\]]+)\]\(([^)]+)\)")


def set_run_font(run) -> None:
    run.font.name = "宋体"
    run.font.size = Pt(11)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")


def configure_styles(doc: Document) -> None:
    style = doc.styles["Normal"]
    style.font.name = "宋体"
    style.font.size = Pt(11)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")

    for table_style_name in ("Table", "Table Grid"):
        if table_style_name not in doc.styles:
            continue
        table_style = doc.styles[table_style_name]
        table_style.font.name = "宋体"
        table_style.font.size = Pt(11)
        table_style.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")

    for i in range(1, 5):
        heading_style = doc.styles[f"Heading {i}"]
        heading_style.font.name = "宋体"
        heading_style.font.size = Pt(11 + max(0, 4 - i))
        heading_style.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")


def normalize_inline(text: str) -> str:
    return LINK_RE.sub(r"\1 (\2)", text)


def add_inline_runs(paragraph, text: str) -> None:
    text = normalize_inline(text)
    parts = text.split("**")
    for index, part in enumerate(parts):
        if not part:
            continue
        run = paragraph.add_run(part)
        run.bold = index % 2 == 1
        set_run_font(run)


def add_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    table = doc.add_table(rows=1, cols=len(rows[0]))
    table.style = "Table" if "Table" in doc.styles else "Table Grid"
    table.autofit = True
    header_cells = table.rows[0].cells
    for idx, value in enumerate(rows[0]):
        header_cells[idx].text = ""
        p = header_cells[idx].paragraphs[0]
        add_inline_runs(p, value)
    for row in rows[1:]:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            cells[idx].text = ""
            p = cells[idx].paragraphs[0]
            add_inline_runs(p, value)
    doc.add_paragraph("")


def parse_table(lines: list[str], start: int) -> tuple[list[list[str]], int]:
    rows: list[list[str]] = []
    index = start
    while index < len(lines) and lines[index].strip().startswith("|"):
        line = lines[index].strip().strip("|")
        cells = [cell.strip() for cell in line.split("|")]
        if cells and all(cell.startswith("---") or cell.startswith(":-") for cell in cells):
            index += 1
            continue
        rows.append(cells)
        index += 1
    return rows, index


def markdown_to_docx(markdown_path: Path, output_path: Path) -> None:
    doc = Document()
    configure_styles(doc)
    section = doc.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    lines = markdown_path.read_text(encoding="utf-8").splitlines()
    index = 0
    while index < len(lines):
        raw = lines[index]
        stripped = raw.strip()

        if not stripped:
            index += 1
            continue

        if stripped.startswith("|"):
            rows, index = parse_table(lines, index)
            add_table(doc, rows)
            continue

        if stripped.startswith("#"):
            level = min(len(stripped) - len(stripped.lstrip("#")), 4)
            text = stripped[level:].strip()
            paragraph = doc.add_heading("", level=level)
            add_inline_runs(paragraph, text)
            index += 1
            continue

        if stripped.startswith("> "):
            paragraph = doc.add_paragraph(style="Normal")
            add_inline_runs(paragraph, stripped[2:].strip())
            index += 1
            continue

        if re.match(r"^\d+\.\s+", stripped):
            paragraph = doc.add_paragraph(style="Normal")
            add_inline_runs(paragraph, stripped)
            index += 1
            continue

        if stripped.startswith("- "):
            paragraph = doc.add_paragraph(style="Normal")
            add_inline_runs(paragraph, f"• {stripped[2:].strip()}")
            index += 1
            continue

        paragraph = doc.add_paragraph(style="Normal")
        add_inline_runs(paragraph, stripped)
        index += 1

    doc.save(output_path)


def main() -> int:
    if len(sys.argv) != 3:
        print("Usage: python3 generate_a_stock_picker_docx.py <input.md> <output.docx>")
        return 1

    markdown_path = Path(sys.argv[1]).resolve()
    output_path = Path(sys.argv[2]).resolve()
    markdown_to_docx(markdown_path, output_path)
    print(output_path)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
